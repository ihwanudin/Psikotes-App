import openpyxl
from docx import Document
from .common import write


def extract(root):
    doc=Document(root/"Konfirm Akhir"/"Lampiran_Butir4_Teks_Item_DASS21_1.docx")
    table=next(t for t in doc.tables if len(t.rows)==22 and t.cell(0,0).text.startswith("No."))
    items=[{"item":int(r.cells[0].text),"scale":r.cells[1].text.split()[0],"text_id":r.cells[2].text} for r in table.rows[1:]]
    wb=openpyxl.load_workbook(root/"Update DASS"/"Bank Narasi Formula HPP Psikotes.xlsx",data_only=True)
    ws=wb["7. DASS21 Cutoff"]
    cutoffs=[{"scale":ws.cell(r,1).value,"category":ws.cell(r,3).value,"level":int(ws.cell(r,4).value),"range_x2":ws.cell(r,5).value} for r in range(2,17)]
    narratives_ws = wb["8. Narasi DASS21"]
    narratives = [{
        "key": narratives_ws.cell(r, 1).value,
        "type": narratives_ws.cell(r, 2).value,
        "subscale": narratives_ws.cell(r, 3).value,
        "category": narratives_ws.cell(r, 4).value,
        "level": int(narratives_ws.cell(r, 5).value),
        "text_id": narratives_ws.cell(r, 6).value,
        "text_jp": narratives_ws.cell(r, 7).value,
        "follow_up_trigger": narratives_ws.cell(r, 8).value,
    } for r in range(2, 22)]
    follow_up = {
        "monitoring": {
            "condition": narratives_ws.cell(23, 2).value,
            "text_id": narratives_ws.cell(23, 6).value,
            "text_jp": narratives_ws.cell(23, 7).value,
        },
        "referral": {
            "condition": narratives_ws.cell(24, 2).value,
            "text_id": narratives_ws.cell(24, 6).value,
            "text_jp": narratives_ws.cell(24, 7).value,
        },
    }
    data={"items":items,"cutoffs":cutoffs,"multiplier":2,
          "narratives":narratives,"follow_up":follow_up}
    write("dass21.json",data)
    return data
