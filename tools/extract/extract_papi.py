import openpyxl
from .common import bounds, write


ZONE_COLUMNS = {
    3: ("yellow_low", "Kuning Ekstrem Rendah"),
    4: ("blue_low", "Biru Moderat Bawah"),
    5: ("white", "Putih Optimal/Adaptif"),
    6: ("blue_high", "Biru Moderat Atas"),
    7: ("yellow_high", "Kuning Ekstrem Tinggi"),
}


def extract(root):
    wb = openpyxl.load_workbook(root / "Master Kamus Tes PAPI Kostick.xlsx", data_only=True)
    ws = wb["Dimensi Mapping"]
    mapping = [{"item":int(ws.cell(r,1).value), "type":ws.cell(r,2).value,
                "a":str(ws.cell(r,3).value).split("→")[-1].strip(), "b":str(ws.cell(r,4).value).split("→")[-1].strip()}
               for r in range(8,98)]
    lookup = openpyxl.load_workbook(root / "PSIKOTEST" / "Skoring" / "Tabel Lookup Skoring Psikotes v1.1.xlsx", data_only=True)
    bands = lookup["08 PAPI Pita 20 Dim"]
    dimensions = {}
    for row in range(6, 26):
        code = bands.cell(row, 1).value
        zones = []
        for column, (zone, label) in ZONE_COLUMNS.items():
            raw_range = bands.cell(row, column).value
            if raw_range is None:
                continue
            lo, hi = bounds(raw_range)
            zones.append({"zone": zone, "label": label, "lo": lo, "hi": hi, "range": raw_range})
        dimensions[code] = {
            "dimension": bands.cell(row, 2).value,
            "zones": zones,
            "hpp_usage": bands.cell(row, 8).value,
        }

    conversion = {bands.cell(row, 1).value: int(bands.cell(row, 2).value) for row in range(31, 35)}
    band_scores = {
        "white": conversion["Putih (Sesuai)"],
        "blue_low": conversion["Biru (Normatif Adaptif / Optimal)"],
        "blue_high": conversion["Biru (Normatif Adaptif / Optimal)"],
        "yellow_low": conversion["Kuning di ujung bawah dimensi"],
        "yellow_high": conversion["Kuning di ujung atas dimensi"],
    }
    white = {code: [zone["lo"], zone["hi"]]
             for code, dimension in dimensions.items()
             for zone in dimension["zones"] if zone["zone"] == "white"}
    data = {"mapping": mapping, "bands": dimensions, "band_scores": band_scores, "white_zones": white}
    write("papi.json", data)
    return data
