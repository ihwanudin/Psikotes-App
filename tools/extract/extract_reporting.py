import openpyxl
from .common import write


def extract(root):
    wb=openpyxl.load_workbook(root/"Update DASS"/"Bank Narasi Formula HPP Psikotes.xlsx",data_only=True)
    ga=wb["3. Grey Area"]
    base_standards={}
    critical=[]
    for r in range(11,29):
        code=ga.cell(r,2).value
        standard=ga.cell(r,5).value
        base_standards[code]=int(standard) if isinstance(standard,(int,float)) else None
        if str(ga.cell(r,6).value).startswith("YA"):
            critical.append(code)
    fields=[]
    for r in range(33,39):
        raised_parts=[part.strip() for part in str(ga.cell(r,6).value).split(";") if "->" in part]
        raised=[part.split()[0] for part in raised_parts]
        raised_standards={part.split()[0]:int(part.rsplit("->",1)[1].strip()) for part in raised_parts}
        interest_text=str(ga.cell(r,5).value)
        interest=interest_text.split()[0] if interest_text!="-" else None
        interest_standard=int(interest_text.rsplit(" ",1)[1]) if interest is not None else None
        fields.append({"code":ga.cell(r,2).value,"required_interest":interest,
                       "required_interest_standard":interest_standard,
                       "raised_to_4":raised,"raised_standards":raised_standards})
    nar=wb["2. Narasi Aspek"]
    narratives=[{"key":nar.cell(r,1).value,"aspect":nar.cell(r,2).value,"level":int(nar.cell(r,5).value),"id":nar.cell(r,10).value,"jp":nar.cell(r,11).value} for r in range(2,92)]
    con=wb["12. Konektor Narasi"]
    connectors=[{"group":con.cell(r,1).value,"order":con.cell(r,3).value,"text":con.cell(r,4).value} for r in range(2,14) if con.cell(r,1).value]
    data={"standard_version":ga.cell(50,2).value,"base_standards":base_standards,
          "fields":fields,"critical":critical,"narratives":narratives,"connectors":connectors}
    write("reporting.json",data); return data
