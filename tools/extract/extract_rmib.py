import openpyxl
from .common import write


def extract(root):
    lookup=openpyxl.load_workbook(root/"PSIKOTEST"/"Skoring"/"Tabel Lookup Skoring Psikotes v1.1.xlsx",data_only=True)
    ws=lookup["10 RMIB Kategori"]
    cats=[{"index":int(ws.cell(r,1).value),"code":ws.cell(r,2).value,"name":ws.cell(r,3).value} for r in range(6,18)]
    rotation=[{"group":g,"position":p,"category":((p+g-2)%12)+1} for g in range(1,10) for p in range(1,13)]
    rank={str(int(lookup["11 RMIB Rank ke Skor"].cell(r,1).value)):int(lookup["11 RMIB Rank ke Skor"].cell(r,2).value) for r in range(5,17)}
    data={"categories":cats,"rotation":rotation,"rank_to_score":rank}; write("rmib.json",data); return data
