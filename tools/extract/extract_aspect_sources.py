import json
import re

import openpyxl

from .common import OUTPUT


VERSION = "ASPECT-SOURCES-2026.09"
ASPECTS = ["A1", "A2", "B1", "B2", "B3", "B4", "C1", "C2", "C3", "C4", "C5", "C6", "C7", "D1", "D2", "D3", "D4", "D5"]
WORKBOOK = "Bank Narasi Formula HPP Psikotes.xlsx"


def build(root):
    workbook = openpyxl.load_workbook(
        root / "Update DASS" / WORKBOOK,
        data_only=True,
        read_only=True,
    )
    try:
        formulas = _formula_sources(workbook["13. Skala Formula"])
        rmib = _rmib_sources(workbook["4. Konversi Skor"])
    finally:
        workbook.close()
    aspects = {}

    for aspect in ASPECTS:
        sources = formulas[aspect]
        if aspect.startswith("D"):
            if sources != ["RMIB"]:
                raise ValueError(f"{aspect} must declare exactly one RMIB source")
            sources = [rmib[aspect]]
        aspects[aspect] = sources

    _validate_contract(aspects)

    return {"version": VERSION, "aspects": aspects}


def extract(root, output=OUTPUT, identifier_root=OUTPUT):
    data = build(root)
    validate_source_identifiers(data, identifier_root)
    output.mkdir(parents=True, exist_ok=True)
    (output / "aspect_sources.json").write_bytes(_serialize(data))

    return data


def validate_source_identifiers(data, output=OUTPUT):
    ist = _load_json(output / "ist.json")
    papi = _load_json(output / "papi.json")
    kraepelin = _load_json(output / "kraepelin.json")
    rmib = _load_json(output / "rmib.json")

    allowed = {"IST_IQ"}
    allowed.update(f"IST_{item['subtest']}" for item in ist["keys"])
    allowed.update(f"IST_{code}" for code in ist["norms"])
    allowed.update(f"PAPI_{code}" for code in papi["bands"])
    allowed.update(f"KRAEPELIN_{band['factor'].upper()}" for band in kraepelin["score_bands"])
    allowed.update(f"RMIB_{category['code']}" for category in rmib["categories"])

    unknown = sorted({
        source
        for sources in data["aspects"].values()
        for source in sources
        if source not in allowed
    })
    if unknown:
        raise ValueError(f"Unknown F0 aspect sources: {', '.join(unknown)}")


def _formula_sources(sheet):
    header_row, columns = _find_header(sheet, ["KODE", "JML SUMBER", "SUMBER TERURAI"])
    formulas = {}

    for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        aspect = _text(row[columns["KODE"]])
        if not re.fullmatch(r"[ABCD]\d", aspect):
            continue
        if aspect in formulas:
            raise ValueError(f"Duplicate aspect formula: {aspect}")

        declared = row[columns["JML SUMBER"]]
        if isinstance(declared, float) and declared.is_integer():
            declared = int(declared)
        if isinstance(declared, str) and declared.isdigit():
            declared = int(declared)
        if not isinstance(declared, int):
            raise ValueError(f"{aspect} has a non-integer source count")

        sources = _parse_formula(_text(row[columns["SUMBER TERURAI"]]))
        if declared != len(sources):
            raise ValueError(f"{aspect} declares {declared} sources but contains {len(sources)}")
        formulas[aspect] = sources

    if list(formulas) != ASPECTS:
        raise ValueError("Aspect formulas must contain exactly A1-D5 in canonical order")

    return formulas


def _rmib_sources(sheet):
    header_row, columns = _find_header(sheet, ["KODE", "DIPAKAI DI HPP", "ASPEK HPP"])
    mapped = {}

    for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        if _text(row[columns["DIPAKAI DI HPP"]]).casefold() != "ya":
            continue
        aspect = _text(row[columns["ASPEK HPP"]])
        if not re.fullmatch(r"D[1-5]", aspect):
            raise ValueError(f"RMIB category maps to an invalid HPP aspect: {aspect}")
        if aspect in mapped:
            raise ValueError(f"Duplicate RMIB mapping: {aspect}")
        code = _text(row[columns["KODE"]])
        mapped[aspect] = f"RMIB_{code}"

    if set(mapped) != set(ASPECTS[-5:]):
        raise ValueError("RMIB mapping must contain exactly D1-D5")

    return mapped


def _parse_formula(value):
    sources = []
    for component in (part.strip() for part in value.split("+")):
        if component == "RMIB":
            sources.append(component)
            continue
        match = re.fullmatch(r"(IST|PAPI|Kraepelin)\s+(.+)", component, flags=re.IGNORECASE)
        if match is None:
            raise ValueError(f"Unknown aspect-source formula component: {component}")
        instrument, measure = match.groups()
        instrument = instrument.upper()
        measure = measure.strip()
        if instrument == "IST" and measure.casefold() == "iq total":
            measure = "IQ"
        if not re.fullmatch(r"[A-Za-z]+", measure):
            raise ValueError(f"Invalid aspect-source measure: {measure}")
        sources.append(f"{instrument}_{measure.upper()}")

    return sources


def _find_header(sheet, required):
    for number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        normalized = [_text(value) for value in row]
        if all(column in normalized for column in required):
            return number, {column: normalized.index(column) for column in required}
    raise ValueError(f"Missing columns in {sheet.title}: {', '.join(required)}")


def _validate_contract(aspects):
    if list(aspects) != ASPECTS:
        raise ValueError("Aspect-source mapping must be ordered A1-D5")
    for aspect, sources in aspects.items():
        if not sources or len(sources) != len(set(sources)):
            raise ValueError(f"{aspect} sources must be non-empty and unique")
    if sum(len(sources) for sources in aspects.values()) != 42:
        raise ValueError("Aspect-source mapping must contain exactly 42 associations")
    if aspects["A1"] != ["IST_IQ"]:
        raise ValueError("A1 must use only the derived IST_IQ source")
    if aspects["D5"] != ["RMIB_S.Se"]:
        raise ValueError("D5 must use only RMIB_S.Se")
    flattened = [source for sources in aspects.values() for source in sources]
    if "RMIB_Prs" in flattened or any(source.startswith("DASS") for source in flattened):
        raise ValueError("Persuasive and DASS are not aspect sources")


def _load_json(path):
    return json.loads(path.read_text(encoding="utf-8"))


def _serialize(data):
    return json.dumps(data, ensure_ascii=False, indent=2).encode("utf-8")


def _text(value):
    return "" if value is None else str(value).strip()
