import hashlib
import json
from pathlib import Path

import openpyxl
from .common import write, OUTPUT

EXPECTED_GRID_HASH = "6df51224c36bea665b9f8b0f8792139489f3c3204476089134657e3ae6ff9ee0"


def extract(root):
    scoring_path = root / "PSIKOTEST" / "Skoring" / "Tabel Lookup Skoring Psikotes v1.1.xlsx"
    digit_grid = extract_digit_grid(root)

    if scoring_path.is_file():
        wb = openpyxl.load_workbook(scoring_path, data_only=True)
        ws = wb["06 KRP Cutoff ke Skor"]
        bands = [{"factor": ws.cell(r, 1).value, "group": ws.cell(r, 2).value, "score": int(ws.cell(r, 3).value),
                  "lo": ws.cell(r, 4).value, "hi": ws.cell(r, 5).value, "category": ws.cell(r, 6).value}
                 for r in range(6, 241)]
        data = {"version": "F2-2026.09", "hanker_formula": "slope_b_x_50",
                "panker_achievement": "correct_plus_incorrect",
                "factor_rounding": {"stage": "before_band_lookup", "precision": 3, "mode": "half_up"},
                "score_bands": bands,
                "digit_grid": digit_grid}
    else:
        existing_path = OUTPUT / "kraepelin.json"
        if not existing_path.is_file():
            raise FileNotFoundError(
                f"File scoring tidak ditemukan di {scoring_path} dan "
                f"tidak ada kraepelin.json yang sudah ada di {existing_path}."
            )
        data = json.loads(existing_path.read_text(encoding="utf-8"))
        data["digit_grid"] = digit_grid

    write("kraepelin.json", data)
    return data


def extract_digit_grid(root):
    wb = openpyxl.load_workbook(
        root / "outputs" / "kraepelin" / "Soal_LJK_Kraepelin_sel_editable.xlsx",
        data_only=True,
    )
    ws = wb["Halaman 2"]

    # Rows 3-30 (28 rows), columns B-AY (columns 2-51 in 1-based indexing = 50 columns).
    # Column A = row labels, row 2 = column labels — neither is data.
    grid = []
    for row_idx in range(3, 31):
        row = []
        for col_idx in range(2, 52):
            value = ws.cell(row_idx, col_idx).value
            if value is None:
                raise ValueError(
                    f"Sel kosong di baris {row_idx} kolom {col_idx} — "
                    "semua sel harus bilangan bulat 1-9."
                )
            row.append(int(value))
        grid.append(row)

    # Integrity gate: hash must match the independently verified value.
    computed = hashlib.sha256(
        json.dumps(grid, separators=(",", ":")).encode()
    ).hexdigest()
    if computed != EXPECTED_GRID_HASH:
        raise SystemExit(
            f"Hash grid TIDAK cocok.\n"
            f"  Diharapkan: {EXPECTED_GRID_HASH}\n"
            f"  Dihitung:   {computed}\n"
            "Ekstraksi dihentikan. Periksa file sumber atau definisi grid."
        )

    return {
        "grid": grid,
        "sha256": computed,
        "numbers_per_column": 28,
        "answer_slots_per_column": 27,
    }
