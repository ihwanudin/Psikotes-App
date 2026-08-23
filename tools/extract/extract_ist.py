import re
import openpyxl
from .common import cell_text, write


def score_bands(sheet, include_split_status=False):
    result = []
    for row in range(6, 16):
        band = {
            "score": int(sheet.cell(row, 1).value),
            "lo": sheet.cell(row, 2).value,
            "hi": sheet.cell(row, 3).value,
            "range": sheet.cell(row, 4).value,
            "category": sheet.cell(row, 5).value,
        }
        if include_split_status:
            band["split_status"] = sheet.cell(row, 6).value
        result.append(band)
    return result


def extract(root):
    wb = openpyxl.load_workbook(root / "Penilaian IST.xlsx", data_only=False)
    ws = wb["Proses"]
    blocks = {"SE": (2, 21), "WA": (22, 41), "AN": (42, 61), "RA": (78, 97),
              "ZR": (98, 117), "FA": (118, 137), "WU": (138, 157), "ME": (158, 177)}
    keys = []
    for code, (lo, hi) in blocks.items():
        for item, col in enumerate(range(lo, hi + 1), 1):
            key_cell = ws.cell(3, col)
            keys.append({"subtest": code, "item": item,
                         "key": cell_text(key_cell.value, key_cell.number_format)})
    ge = []
    for item, col in enumerate(range(62, 78), 1):
        formula = ws.cell(5, col).value
        formula = formula.text if hasattr(formula, "text") else str(formula)
        entries = [{"answer": a, "score": int(s)} for a, s in re.findall(r'="([^"]+)",(1|2)', formula)]
        ge.append({"item": item, "answers": entries})
    lookup = openpyxl.load_workbook(root / "PSIKOTEST" / "Skoring" / "Tabel Lookup Skoring Psikotes v1.1.xlsx", data_only=True)
    ns = lookup["01 IST RW ke SW"]
    norms = {code: {} for code in ["SE","WA","AN","RA","ZR","FA","WU","ME","GE"]}
    for row in range(6, 39):
        rw = ns.cell(row, 1).value
        for col, code in enumerate(["SE","WA","AN","RA","ZR","FA","WU","ME","GE"], 2):
            value = ns.cell(row, col).value
            if value is not None: norms[code][str(int(rw))] = int(value)
    iq = [{"lo": int(ns2.cell(r,1).value), "hi": int(ns2.cell(r,2).value), "iq": int(ns2.cell(r,3).value)}
          for ns2 in [lookup["03 IST RWtotal ke IQ"]] for r in range(6, 62)]
    sw_score_bands = score_bands(lookup["02 IST SW ke Skor"], include_split_status=True)
    iq_score_bands = score_bands(lookup["04 IST IQ ke Skor"])
    data = {"version":"F0-2026.08", "keys":keys, "ge_dictionary":ge, "norms":norms, "iq_ranges":iq,
            "sw_score_bands":sw_score_bands, "iq_score_bands":iq_score_bands,
            "match_strategy":{"SE":"option_prefix","WA":"option_prefix","AN":"option_prefix",
                              "RA":"formula_character_membership","ZR":"formula_character_membership",
                              "FA":"exact","WU":"exact","ME":"option_prefix"}}
    write("ist.json", data); return data
